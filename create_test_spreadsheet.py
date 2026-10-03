import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
RESULTS_DIR = BASE_DIR / "results"
JSON_FILE = RESULTS_DIR / "agent_test_results.json"
EXCEL_FILE = RESULTS_DIR / "agent_test_cases.xlsx"

with open(JSON_FILE, "r", encoding="utf-8") as file:
    data = json.load(file)

tests = data["tests"]

workbook = Workbook()

sheet = workbook.active
sheet.title = "Test Cases"

headers = [
    "Test ID",
    "Test Name",
    "User Request",
    "Data / Tools Used",
    "Expected Response / Behavior",
    "Actual Agent Response",
    "Tool Calls Made",
    "Pass / Fail",
    "Error Analysis"
]

sheet.append(headers)

for cell in sheet[1]:
    cell.font = Font(bold=True)
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )

for test in tests:
    sheet.append([
        test["test_id"],
        test["test_name"],
        test["user_request"],
        test["data_tools_used"],
        test["expected_behavior"],
        test["actual_response"],
        "\n".join(test["tool_calls"]),
        test["status"],
        test["error_analysis"]
    ])

for row in sheet.iter_rows():
    for cell in row:
        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True
        )

widths = {
    "A": 12,
    "B": 28,
    "C": 45,
    "D": 45,
    "E": 55,
    "F": 70,
    "G": 50,
    "H": 15,
    "I": 60
}

for column, width in widths.items():
    sheet.column_dimensions[column].width = width

sheet.freeze_panes = "A2"
sheet.auto_filter.ref = sheet.dimensions

summary = workbook.create_sheet("Summary")

summary_data = [
    ["AI Agent Testing Summary", ""],
    ["Generated At", data["generated_at"]],
    ["Total Tests", data["total_tests"]],
    ["Passed", data["passed"]],
    ["Failed", data["failed"]],
    ["Pass Rate", f"{(data['passed'] / data['total_tests']) * 100:.2f}%"],
    ["", ""],
    ["Error / Edge Case Coverage", ""],
    ["Invalid Tool Arguments", "T03"],
    ["Missing Data", "T05"],
    ["Ambiguous Request", "T06"],
    ["Unauthorized Action", "T07"],
    ["Conflicting / Changing Data", "T10"],
    ["High-Level Agent Objective", "T11"],
    ["End-to-End Analysis", "T12"],
    ["", ""],
    ["Important Note", "Some edge cases are controlled validation tests. They verify that the system handles unsafe, invalid, missing, or ambiguous situations correctly; they do not claim that the LLM independently made the corresponding mistake."]
]

for row in summary_data:
    summary.append(row)

for cell in summary["A"]:
    cell.font = Font(bold=True)

summary.column_dimensions["A"].width = 32
summary.column_dimensions["B"].width = 100

for row in summary.iter_rows():
    for cell in row:
        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True
        )

workbook.save(EXCEL_FILE)

print("=" * 70)
print("EXCEL TEST REPORT CREATED")
print("=" * 70)
print(f"File: {EXCEL_FILE}")
print(f"Total Tests: {data['total_tests']}")
print(f"Passed: {data['passed']}")
print(f"Failed: {data['failed']}")
print("=" * 70)